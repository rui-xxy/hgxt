"""Extract the three non-packaging sheets of a Fenglian monthly report.

Usage: python extract-fenglian.py /path/to/丰联生产日报表2026.9月.xlsx
The schema and dated snapshot are written beside the Prisma form schemas/data.
Excel formula caches are read as displayed; this script never recalculates the source.
"""

import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


ROOT = Path(__file__).resolve().parents[1]
FIELDS = []


def add(sheet, column, field_id, title, group, subgroup, unit='t'):
    FIELDS.append((sheet, column, field_id, title, group, subgroup, unit))


def triplet(sheet, columns, ids, label, group, subgroup, unit='t', verbs=('购入', '耗用', '库存')):
    for column, field_id, verb in zip(columns, ids, verbs, strict=True):
        add(sheet, column, field_id, f'{label}{verb}', group, subgroup, unit)


THREE = '三车间'
STANDARD = '标准厂房'
PRODUCTS = '标准厂房产成品'
MATERIALS = '标准厂房原辅料'
UTILITY = '公共'

triplet(THREE, ('B', 'C', 'D'), ('field_201', 'field_202', 'field_203'),
        '焦磷酸哌嗪湿料', THREE, '湿料', verbs=('产量', '出库', '库存'))
for columns, ids, label, unit in [
    (('E', 'F', 'G'), ('field_004', 'field_005', 'field_006'), '85%磷酸', 't'),
    (('H', 'I', 'J'), ('field_007', 'field_008', 'field_009'), '68%哌嗪', 't'),
    (('K', 'L', 'M'), ('field_010', 'field_011', 'field_012'), '椰壳活性炭', 'kg'),
]:
    triplet(THREE, columns, ids, label, THREE, label, unit)

for column, field_id, title in [
    ('B', 'field_dry_drying', '焦磷酸哌嗪干燥量'),
    ('C', 'field_dry_crushing', '焦磷酸哌嗪粉碎量'),
    ('D', 'field_204', '焦磷酸哌嗪打包数'),
    ('E', 'field_205', '焦磷酸哌嗪销量'),
    ('F', 'field_dry_self_use', '焦磷酸哌嗪自耗数'),
    ('G', 'field_dry_sluggish_output', '焦磷酸哌嗪呆滞品生产数'),
    ('H', 'field_dry_sluggish_stock', '焦磷酸哌嗪呆滞品库存'),
    ('I', 'field_dry_returns', '焦磷酸哌嗪退货数'),
    ('J', 'field_206', '焦磷酸哌嗪成品库存'),
]:
    add('标准厂房', column, field_id, title, PRODUCTS, '焦磷酸哌嗪干料')

for columns, key, name, extra in [
    (('K', 'L', 'M'), '3500_high', '3500阻燃母粒（优）', False),
    (('N', 'O', 'P'), '3500_standard', '3500阻燃母粒（普）', False),
    (('Q', 'R', 'S'), '3500_new', '新3500阻燃母粒', False),
    (('T', 'U', 'V'), '3500_rework', '返工3500阻燃母粒', False),
    (('W', 'X', 'Y'), '4000_masterbatch', '4000阻燃母粒', False),
    (('Z', 'AA', 'AB'), '4500_high', '4500阻燃母粒（优）', False),
    (('AC', 'AD', 'AE'), '4500_standard', '4500阻燃母粒（普）', False),
    (('AG', 'AH', 'AI'), '3000_halogen_free', '3000无卤阻燃剂', False),
    (('AJ', 'AK', 'AM'), '4000_halogen_free', '4000无卤阻燃剂', True),
    (('AN', 'AO', 'AP'), '3500_halogen_free_1', '3500无卤阻燃剂（第一组）', False),
    (('AQ', 'AR', 'AS'), '3500_halogen_free_2', '3500无卤阻燃剂（第二组）', False),
    (('AT', 'AU', 'AV'), 'mb5000', 'MB-5000', False),
]:
    if extra:
        for column, verb, title in zip(('AJ', 'AK', 'AL', 'AM'),
                                       ('output', 'sales', 'self_use', 'stock'),
                                       ('产量', '销量', '自耗数', '库存'), strict=True):
            add('标准厂房', column, f'field_{key}_{verb}', name + title, PRODUCTS, name)
    else:
        triplet('标准厂房', columns,
                tuple(f'field_{key}_{verb}' for verb in ('output', 'sales', 'stock')),
                name, PRODUCTS, name, verbs=('产量', '销量', '库存'))
add('标准厂房', 'AF', 'field_4500_returns', '4500阻燃母粒退货数', PRODUCTS, '4500阻燃母粒（普）')

raw = [
    ('AX', '035', '焦磷酸哌嗪湿料', '入库'),
    ('BA', '038', 'MPP母粒级', '购入'),
    ('BD', '041', 'MPP普通级', '购入'),
    ('BG', '044', 'MPP国标级', '购入'),
    ('BJ', '047', 'PP450粉', '购入'),
    ('BM', '053', '氧化锌', '购入'),
    ('BP', '056', '2乙基次磷酸铝', '购入'),
    ('BS', '059', '三聚氰胺氰尿酸盐MCA', '购入'),
    ('BV', '062', 'PETS季戊四醇硬脂酸脂', '购入'),
    ('BY', '065', '抗氧化剂168', '购入'),
    ('CB', '068', '抗氧化剂1010', '购入'),
    ('CE', '071', '气相二氧化硅', '购入'),
    ('CH', '074', '助剂C', '购入'),
    ('CK', '080', '化学试剂助剂（改制剂810）', '购入'),
    ('CN', '120', '3-氨丙基三乙氧基硅烷KH550', '购入'),
    ('CQ', None, 'PP8285E', '购入'),
    ('CT', None, 'DMPY', '购入'),
]
from openpyxl.utils import get_column_letter

for start, first_id, label, incoming in raw:
    first = column_index_from_string(start)
    cols = tuple(get_column_letter(first + offset) for offset in range(3))
    ids = (tuple(f'field_{int(first_id) + offset:03d}' for offset in range(3))
           if first_id else tuple(f'field_{label.lower()}_{verb}' for verb in ('purchase', 'consumption', 'stock')))
    triplet('标准厂房', cols, ids, label, MATERIALS, label,
            verbs=(incoming, '耗用', '库存'))

for column, field_id, title, subgroup, unit in [
    ('B', 'field_nitrogen_purchase', '液氮购入', '液氮', 't'),
    ('C', 'field_nitrogen_consumption', '液氮耗用', '液氮', 't'),
    ('D', 'field_nitrogen_stock', '液氮剩余', '液氮', 't'),
    ('E', 'field_water_cumulative', '水表累计读数', '自来水', ''),
    ('F', 'field_electricity_cumulative', '电表累计读数', '用电', ''),
    ('G', 'field_gas_purchase', '天然气购入', '天然气', 'm³'),
    ('H', 'field_gas_consumption', '天然气耗用', '天然气', 'm³'),
    ('I', 'field_gas_stock', '天然气剩余', '天然气', 'm³'),
]:
    add('公共', column, field_id, title, UTILITY, subgroup, unit)


def as_number(value, sheet, coordinate):
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f'{sheet}!{coordinate} 为布尔值')
    if isinstance(value, (int, float)):
        result = value
    elif isinstance(value, str) and re.fullmatch(r'\d+(?:\.\d+)?', value):
        result = float(value)
    else:
        raise ValueError(f'{sheet}!{coordinate} 不是数字：{value!r}')
    return result


def main(path):
    workbook = load_workbook(path, data_only=True, read_only=False)
    expected = ('三车间', '标准厂房', '公共')
    if any(name not in workbook.sheetnames for name in expected):
        raise ValueError(f'缺少工作表：{expected}')
    field_ids = [entry[2] for entry in FIELDS]
    if len(field_ids) != len(set(field_ids)):
        raise ValueError('字段 ID 重复')
    columns = {sheet: set() for sheet in expected}
    for sheet, column, *_ in FIELDS:
        if column in columns[sheet]:
            raise ValueError(f'{sheet}!{column} 被映射多次')
        columns[sheet].add(column)
    # AW is a merged blank divider between finished goods and raw materials.
    if [len(columns[name]) for name in expected] != [12, 98, 8]:
        raise ValueError(f'生产/原料列映射不完整：{ {name: len(columns[name]) for name in expected} }')
    schema = [{'id': 'field_date', 'title': '日期', 'type': 'date', 'group': THREE,
               'hidden': True, 'required': True}]
    for _, _, field_id, title, group, subgroup, unit in FIELDS:
        main_group = THREE if group == THREE else STANDARD
        section = ('产成品' if subgroup == '湿料' else '原辅料') if group == THREE else (
            '产成品' if group == PRODUCTS else '原辅料' if group == MATERIALS else '水电气')
        item = {'id': field_id, 'title': title, 'type': 'number', 'group': main_group,
                'section': section, 'subgroup': subgroup, 'step': 0.001, 'precision': 3}
        if unit:
            item['unit'] = unit
        schema.append(item)
    rows = {}
    for name in expected:
        sheet = workbook[name]
        entries = []
        for cells in sheet:
            date = cells[0].value
            if not isinstance(date, datetime) or date.year != 2026 or date.month != 9:
                continue
            values = {}
            for source, column, field_id, *_ in FIELDS:
                if source != name:
                    continue
                cell = sheet[f'{column}{cells[0].row}']
                value = as_number(cell.value, name, cell.coordinate)
                if value is not None:
                    values[field_id] = value
            entries.append({'date': date.strftime('%Y-%m-%d'), 'sourceRow': cells[0].row,
                            'values': values})
        if len(entries) != 30 or len({entry['date'] for entry in entries}) != 30:
            raise ValueError(f'{name} 的 9 月日期不完整')
        rows[name] = entries
    (ROOT / 'form-schemas' / 'fenglian.json').write_text(
        json.dumps(schema, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (ROOT / 'data' / 'fenglian-2026-09.json').write_text(
        json.dumps({'from': '2026-09-01', 'to': '2026-09-30', 'sheets': rows},
                   ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    print(f'丰联：{len(schema)} 个字段、每张表 30 天，已生成 schema 和 9 月快照')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit('用法：python extract-fenglian.py <xlsx路径>')
    main(sys.argv[1])
